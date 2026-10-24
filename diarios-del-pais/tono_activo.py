"""Aprendizaje activo para el modelo de tono: elige los titulares donde el modelo duda y suma las etiquetas a mano.
Uso:
  python tono_activo.py seleccionar [--n 60] [--por-medio 4]   -> muestra_activa.xlsx + muestra_activa.json
  python tono_activo.py ingestar muestra_activa.xlsx            -> agrega a labels_activo.json (acumulativo)
  python tono_bert_eval.py                                      -> reentrena con las etiquetas nuevas y compara antes/despues

Incertidumbre: MARGEN entre las dos clases mas probables del modelo combinado (rasgos de nlp.feats + probas XLM-R,
el mismo de tono_bert_eval.py). Se eligen los de menor margen: son los que estan sobre la frontera de decision, donde
una etiqueta mas mueve el modelo. Con 3 clases es casi lo mismo que la entropia, pero el margen ignora cuanto
se reparte la tercera clase (que no cambia la prediccion); la entropia se guarda igual en el json para comparar."""
import json, os, sys, sqlite3, argparse
from collections import Counter
import numpy as np
from sklearn.linear_model import LogisticRegression
from sklearn.feature_extraction import DictVectorizer
from scipy.sparse import hstack, csr_matrix
import nlp

AQUI = os.path.dirname(os.path.abspath(__file__))
CL = ["n", "0", "p"]
NOMBRE = {"n": "Negativo", "0": "Neutro", "p": "Positivo"}
ACTIVO = os.path.join(AQUI, "labels_activo.json")

def _j(f): return json.load(open(os.path.join(AQUI, f), encoding="utf-8"))

# ---------- datos y rasgos (compartidos con tono_bert_eval.py) ----------
def etiquetadas(activo=True):
    """(urls, titulos, tonos, es_original): los 340 originales (muestra_300 + tuyas_40) y, si hay, labels_activo.json"""
    lb, muestra, tuyas = _j("labels.json"), _j("muestra_300.json"), _j("tuyas_40.json")
    U = [x["u"] for x in muestra] + list(tuyas); T = [x["t"] for x in muestra] + [t for t, _ in tuyas.values()]
    Y = [lb[x["u"]][1] for x in muestra] + [s for _, s in tuyas.values()]
    orig = [True] * len(U)
    if activo and os.path.exists(ACTIVO):
        ya = set(U)
        for u, (t, s) in json.load(open(ACTIVO, encoding="utf-8")).items():
            if u not in ya: U.append(u); T.append(t); Y.append(s); orig.append(False)
    return U, T, np.array(Y), np.array(orig)

def rasgos(T, PB, dv=None):
    """matriz del modelo combinado: bolsa de palabras/bigramas/diccionario (nlp.feats) + 3 x probas XLM-R"""
    F = [nlp.feats(t) for t in T]
    if dv is None: dv = DictVectorizer(); X = dv.fit_transform(F)
    else: X = dv.transform(F)
    return hstack([X, csr_matrix(np.asarray(PB) * 3)]).tocsr(), dv

def modelo(): return LogisticRegression(C=3, max_iter=2000, class_weight="balanced")

def _probas_bert(T):
    import bert
    if not bert.disponible(): sys.exit("hace falta BERT (torch + transformers) para las probabilidades de los titulares etiquetados")
    return bert.probas(T)

# ---------- casi duplicados ----------
def _tok(t): return set(nlp._toks(t))
def _parecido(a, b): return len(a & b) / max(1, len(a | b))      # Jaccard de palabras normalizadas
SIM = 0.5

# ---------- seleccionar ----------
def seleccionar(n=60, por_medio=4):
    U, T, Y, _ = etiquetadas()
    print("etiquetados:", len(T), dict(Counter(Y.tolist())))
    print("probabilidades XLM-R de los etiquetados...")
    X, dv = rasgos(T, _probas_bert(T))
    lr = modelo().fit(X, Y)
    db = sqlite3.connect(os.path.join(os.environ.get("DATA_DIR", AQUI), "historico.db"))
    filas = db.execute("SELECT url, titulo, medio, bert_t FROM items WHERE titulo IS NOT NULL").fetchall()
    ya = set(U); vistos = [_tok(t) for t in T]
    cand = [(u, t, m, b) for u, t, m, b in filas if u not in ya and len(_tok(t)) >= 4]
    falta = [i for i, c in enumerate(cand) if not c[3]]
    if falta:                                                    # titulares sin bert_t: se calculan
        import bert
        pb = _probas_bert([cand[i][1] for i in falta])
        for i, p in zip(falta, pb): cand[i] = cand[i][:3] + (bert.a_texto(p),)
    PB = np.array([[float(x) for x in c[3].split(",")] for c in cand])
    P = lr.predict_proba(rasgos([c[1] for c in cand], PB, dv)[0])
    orden = np.argsort(-P, 1); idx = np.arange(len(P))
    margen = P[idx, orden[:, 0]] - P[idx, orden[:, 1]]
    entropia = -(P * np.log(P + 1e-12)).sum(1)
    k = {c: list(lr.classes_).index(c) for c in CL}
    sel, pm, toks_sel = [], Counter(), []
    for i in np.argsort(margen):
        u, t, m, _ = cand[i]; tk = _tok(t)
        if pm[m] >= por_medio: continue
        if any(_parecido(tk, o) >= SIM for o in toks_sel) or any(_parecido(tk, o) >= SIM for o in vistos): continue
        sel.append({"url": u, "titulo": t, "medio": m, "prediccion": str(lr.classes_[orden[i, 0]]),
                    "probabilidades": {c: round(float(P[i, k[c]]), 4) for c in CL},
                    "margen": round(float(margen[i]), 4), "entropia": round(float(entropia[i]), 4)})
        pm[m] += 1; toks_sel.append(tk)
        if len(sel) == n: break
    print(f"candidatos: {len(cand)} | elegidos: {len(sel)} de {len(pm)} medios | margen medio {np.mean([s['margen'] for s in sel]):.3f}"
          f" (en toda la base {margen.mean():.3f}) | predicciones {dict(Counter(s['prediccion'] for s in sel))}")
    json.dump(sel, open(os.path.join(AQUI, "muestra_activa.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    planilla(sel, os.path.join(AQUI, "muestra_activa.xlsx"))

def planilla(sel, salida):
    """mismo estilo que crear_planilla_validacion.py. La prediccion del modelo NO va en la planilla (solo en el json):
    verla ancla al que etiqueta, y en los casos dudosos es justo donde mas sesgaria"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule
    wb = Workbook(); ws = wb.active; ws.title = "Muestra"
    borde = Border(bottom=Side(style="thin", color="D9D9D9"))
    ws.append(["N°", "Medio", "Titular", "Tu tono", "Comentario (opcional)", "Nota", "url"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", size=11); c.fill = PatternFill("solid", fgColor="1F5F8B")
        c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
    ws.row_dimensions[1].height = 26
    for i, s in enumerate(sel, 1):
        ws.append([i, s["medio"], s["titulo"], None, None, "abrir", s["url"]])
        ws.cell(i + 1, 6).hyperlink = s["url"]; ws.cell(i + 1, 6).font = Font(color="0563C1", underline="single")
    for col, w in {"A": 6, "B": 20, "C": 90, "D": 16, "E": 36, "F": 8}.items(): ws.column_dimensions[col].width = w
    ws.column_dimensions["G"].hidden = True                      # clave para ingestar
    ult = ws.max_row
    for r in range(2, ult + 1):
        ws.row_dimensions[r].height = 34
        for c in ws[r][:5]: c.alignment = Alignment(vertical="center", wrap_text=True); c.border = borde
        ws.cell(r, 1).alignment = Alignment(vertical="center", horizontal="center")
        ws.cell(r, 2).font = Font(color="666666")
        ws.cell(r, 4).alignment = Alignment(vertical="center", horizontal="center"); ws.cell(r, 4).font = Font(bold=True)
        ws.cell(r, 4).fill = PatternFill("solid", fgColor="FFF8DC")
    dv = DataValidation(type="list", formula1='"Negativo,Neutro,Positivo"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Valor no válido", error="Elegí Negativo, Neutro o Positivo de la lista.")
    dv.promptTitle = "Tu tono"; dv.prompt = "Elegí de la lista"; dv.showInputMessage = True
    ws.add_data_validation(dv); dv.add(f"D2:D{ult}")
    for v, col in (("Negativo", "F4C7C3"), ("Neutro", "E3E5E1"), ("Positivo", "C9E7D2")):
        ws.conditional_formatting.add(f"D2:D{ult}", CellIsRule(operator="equal", formula=[f'"{v}"'], fill=PatternFill("solid", bgColor=col)))
    ws.freeze_panes = "A2"; ws.sheet_view.zoomScale = 110; ws.auto_filter.ref = f"A1:F{ult}"
    ws["I1"] = "Completadas"; ws["I1"].font = Font(bold=True)
    ws["I2"] = f'=COUNTA(D2:D{ult})&" de "&COUNTA(C2:C{ult})'; ws["I2"].font = Font(size=14, bold=True, color="1F5F8B")
    ws.column_dimensions["I"].width = 16
    ins = wb.create_sheet("Instrucciones", 0)
    texto = [("Cómo completar (aprendizaje activo, titulares difíciles)", True), ("", False),
        ("Estos titulares los eligió el modelo porque son los que más dudas le generan: es normal que cueste decidir.", False),
        ("Por eso no se muestra lo que predijo el modelo: verlo te influiría justo en los casos donde más importa tu criterio.", False), ("", False),
        ("1. Andá a la hoja «Muestra».", False),
        ("2. Leé cada titular y elegí en «Tu tono» una opción de la lista: Negativo, Neutro o Positivo.", False),
        ("3. Criterio: ¿el hecho que cuenta el titular le hace bien o mal a las personas involucradas?", False),
        ("      Negativo: un crimen, una crisis, un conflicto, una muerte, un aumento de precios.", False),
        ("      Neutro: información de servicio o agenda («hora y canal del partido», «cómo recuperar un archivo»).", False),
        ("      Positivo: un logro, una celebración, una mejora, algo que se espera con gusto.", False),
        ("4. Se mide el hecho, no la postura del medio ni si estás de acuerdo con lo que pasó.", False),
        ("5. Si dudás, elegí lo primero que se te ocurra y anotá la duda en «Comentario». Podés dejar filas vacías: no se usan.", False),
        ("6. «abrir» lleva a la nota por si el titular solo no alcanza. No borres ni reordenes filas.", False),
        ("7. Cuando termines: Ctrl+S, cerrá Excel y corré:  python tono_activo.py ingestar muestra_activa.xlsx", False)]
    for i, (t, b) in enumerate(texto, 1): ins.cell(i, 1, t).font = Font(bold=b, size=14 if b else 11)
    ins.column_dimensions["A"].width = 120; ins.sheet_view.showGridLines = False
    wb.save(salida); print("creado", salida, "-", ult - 1, "titulares")

# ---------- ingestar ----------
VAL = {"negativo": "n", "neutro": "0", "positivo": "p", "n": "n", "0": "0", "p": "p", "neg": "n", "neu": "0", "pos": "p"}
def ingestar(ruta):
    from openpyxl import load_workbook
    from sklearn.metrics import confusion_matrix
    ws = load_workbook(ruta, data_only=True)["Muestra"]
    cab = [c.value for c in ws[1]]; iu, it, iy = cab.index("url"), cab.index("Titular"), cab.index("Tu tono")
    pj = os.path.join(os.path.dirname(os.path.abspath(ruta)), "muestra_activa.json")
    meta = {s["url"]: s for s in json.load(open(pj if os.path.exists(pj) else os.path.join(AQUI, "muestra_activa.json"), encoding="utf-8"))}
    U0 = set(etiquetadas(activo=False)[0])
    nuevas, malas, vacias = {}, [], 0
    for r, f in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not f or not f[iu]: continue
        v = str(f[iy] or "").strip().lower()
        if not v: vacias += 1; continue
        if v not in VAL: malas.append((r, f[iy])); continue
        if f[iu] in U0: print(f"  fila {r}: ya estaba en los 340 originales, se ignora"); continue
        nuevas[f[iu]] = [f[it], VAL[v]]
    for r, v in malas: print(f"  fila {r}: valor no válido {v!r} (se esperaba Negativo/Neutro/Positivo), se ignora")
    if not nuevas: sys.exit("no hay etiquetas para agregar")
    prev = json.load(open(ACTIVO, encoding="utf-8")) if os.path.exists(ACTIVO) else {}
    rep = sum(1 for u in nuevas if u in prev)
    prev.update(nuevas)
    json.dump(prev, open(ACTIVO, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"etiquetadas: {len(nuevas)} (vacías {vacias}, inválidas {len(malas)}, ya estaban {rep}) | labels_activo.json: {len(prev)} en total")
    print("humano:", {NOMBRE[c]: n for c, n in Counter(s for _, s in nuevas.values()).items()})
    con = [(s, meta[u]["prediccion"]) for u, (_, s) in nuevas.items() if u in meta]
    if con:
        h, m = zip(*con)
        print(f"acuerdo modelo vs humano en estos casos difíciles: {np.mean(np.array(h) == np.array(m)):.0%} de {len(con)}"
              " (se espera bajo: se eligieron por estar en la frontera)")
        M = confusion_matrix(h, m, labels=CL)
        print("matriz de confusión (filas = humano, columnas = modelo)\n          " + "".join(f"{NOMBRE[c]:>10}" for c in CL))
        for c, fila in zip(CL, M): print(f"{NOMBRE[c]:>10}" + "".join(f"{x:>10}" for x in fila))
    print("siguiente paso: python tono_bert_eval.py (reentrena y compara antes/después)")

if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sp = ap.add_subparsers(dest="modo", required=True)
    s = sp.add_parser("seleccionar"); s.add_argument("--n", type=int, default=60); s.add_argument("--por-medio", type=int, default=4)
    i = sp.add_parser("ingestar"); i.add_argument("planilla")
    a = ap.parse_args()
    seleccionar(a.n, a.por_medio) if a.modo == "seleccionar" else ingestar(a.planilla)

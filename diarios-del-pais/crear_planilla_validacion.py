"""Arma muestra_para_validar.xlsx a partir de muestra_para_validar.csv, lista para completar en Excel."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

AQUI = os.path.dirname(os.path.abspath(__file__))
filas = list(csv.reader(open(os.path.join(AQUI, "muestra_para_validar.csv"), encoding="utf-8-sig")))[1:]

wb = Workbook()
ws = wb.active
ws.title = "Muestra"
borde = Border(bottom=Side(style="thin", color="D9D9D9"))
cab = ["N°", "Medio", "Titular", "Tu tono", "Comentario (opcional)"]
ws.append(cab)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor="1F5F8B")
    c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
ws.row_dimensions[1].height = 26

for n, medio, titulo, *_ in filas:
    ws.append([int(n), medio, titulo, None, None])

anchos = {"A": 6, "B": 20, "C": 90, "D": 16, "E": 36}
for col, w in anchos.items(): ws.column_dimensions[col].width = w
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 34
    for c in ws[r]:
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = borde
    ws.cell(r, 1).alignment = Alignment(vertical="center", horizontal="center")
    ws.cell(r, 2).font = Font(color="666666")
    ws.cell(r, 4).alignment = Alignment(vertical="center", horizontal="center")
    ws.cell(r, 4).font = Font(bold=True)
    ws.cell(r, 4).fill = PatternFill("solid", fgColor="FFF8DC")   # celda para completar

ult = ws.max_row
dv = DataValidation(type="list", formula1='"Negativo,Neutro,Positivo"', allow_blank=True, showErrorMessage=True,
                    errorTitle="Valor no válido", error="Elegí Negativo, Neutro o Positivo de la lista.")
dv.promptTitle = "Tu tono"; dv.prompt = "Elegí de la lista"; dv.showInputMessage = True
ws.add_data_validation(dv); dv.add(f"D2:D{ult}")

rojo = PatternFill("solid", bgColor="F4C7C3"); gris = PatternFill("solid", bgColor="E3E5E1"); verde = PatternFill("solid", bgColor="C9E7D2")
ws.conditional_formatting.add(f"D2:D{ult}", CellIsRule(operator="equal", formula=['"Negativo"'], fill=rojo))
ws.conditional_formatting.add(f"D2:D{ult}", CellIsRule(operator="equal", formula=['"Neutro"'], fill=gris))
ws.conditional_formatting.add(f"D2:D{ult}", CellIsRule(operator="equal", formula=['"Positivo"'], fill=verde))

ws.freeze_panes = "A2"
ws.sheet_view.zoomScale = 110
ws.auto_filter.ref = f"A1:E{ult}"

# contador de avance
ws["G1"] = "Completadas"; ws["G1"].font = Font(bold=True)
ws["G2"] = f'=COUNTA(D2:D{ult})&" de "&COUNTA(C2:C{ult})'
ws["G2"].font = Font(size=14, bold=True, color="1F5F8B")
ws.column_dimensions["G"].width = 16

# hoja de instrucciones
ins = wb.create_sheet("Instrucciones", 0)
texto = [
    ("Cómo completar", True),
    ("", False),
    ("1. Andá a la hoja «Muestra».", False),
    ("2. Leé cada titular y elegí en la columna «Tu tono» una opción de la lista: Negativo, Neutro o Positivo.", False),
    ("3. Criterio: ¿el hecho que cuenta el titular le hace bien o mal a las personas involucradas?", False),
    ("      Negativo: un crimen, una crisis, un conflicto, una muerte, un aumento de precios.", False),
    ("      Neutro: información de servicio o agenda («hora y canal del partido», «cómo recuperar un archivo»).", False),
    ("      Positivo: un logro, una celebración, una mejora, algo que se espera con gusto.", False),
    ("4. Clasificá por lo que dice el titular, no por si estás de acuerdo con el hecho.", False),
    ("5. Si dudás, elegí lo primero que se te ocurra. También podés anotar la duda en «Comentario».", False),
    ("6. Cuando termines: Ctrl+S para guardar, cerrá Excel y avisame.", False),
]
for i, (t, b) in enumerate(texto, 1):
    ins.cell(i, 1, t).font = Font(bold=b, size=14 if b else 11)
ins.column_dimensions["A"].width = 120
ins.sheet_view.showGridLines = False

salida = os.path.join(AQUI, "muestra_para_validar.xlsx")
wb.save(salida)
print("creado", salida, "-", ult - 1, "titulares")
